"""Сервис экспорта отчётов.

Реализует F-6.4: экспорт данных в Excel, CSV и PDF.
Использует openpyxl (Excel), csv (стандартная библиотека),
reportlab (PDF).
"""

from datetime import date
from typing import Dict, List, Optional

# Excel
from openpyxl import Workbook

# PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class ExportService:
    """Сервис экспорта отчётов в разные форматы.

    Поддерживает:
    - Excel (.xlsx) — через openpyxl;
    - CSV (.csv) — через стандартную библиотеку;
    - PDF (.pdf) — через reportlab.
    """

    def export_finance_report(
        self,
        path: str,
        format: str,
        period_start: Optional[date] = None,
        period_end: Optional[date] = None,
    ) -> bool:
        """Экспорт финансового отчёта.

        Args:
            path: Путь для сохранения файла.
            format: Формат (excel / csv / pdf).
            period_start: Начало периода.
            period_end: Конец периода.

        Returns:
            True если экспорт успешен.
        """
        # Получаем данные
        data = self._collect_finance_data(period_start, period_end)

        if format == "excel":
            return self._export_to_excel(data, path, period_start, period_end)
        elif format == "csv":
            return self._export_to_csv(data, path)
        elif format == "pdf":
            return self._export_to_pdf(data, path, period_start, period_end)
        return False

    def _collect_finance_data(
        self,
        period_start: Optional[date],
        period_end: Optional[date],
    ) -> Dict:
        """Собирает данные для отчёта.

        Returns:
            Словарь с данными: KPI, платежи, способы оплаты.
        """
        from src.repositories.payment_repository import PaymentRepository
        from src.repositories.order_repository import OrderRepository
        from src.services.report_service import ReportService

        report = ReportService()
        payment_repo = PaymentRepository()

        # Период по умолчанию — последние 30 дней
        if not period_start or not period_end:
            from datetime import timedelta
            period_end = date.today()
            period_start = period_end - timedelta(days=30)

        # Данные по выручке
        revenue = report.get_revenue_for_period(period_start, period_end)

        # Способы оплаты
        payment_methods = report.get_payment_methods_stats(period_start, period_end)

        # Платежи за период
        payments = payment_repo.get_by_date_range(period_start, period_end)

        return {
            "revenue": revenue,
            "payment_methods": payment_methods,
            "payments": payments,
            "period_start": period_start,
            "period_end": period_end,
        }

    # ============================================================
    # Excel
    # ============================================================
    def _export_to_excel(
        self,
        data: Dict,
        path: str,
        period_start: date,
        period_end: date,
    ) -> bool:
        """Экспорт в Excel (.xlsx).

        Args:
            data: Данные отчёта.
            path: Путь к файлу.
            period_start: Начало периода.
            period_end: Конец периода.

        Returns:
            True если успешно.
        """
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Финансовый отчёт"

            # Заголовок
            ws["A1"] = "Финансовый отчёт CleanOps"
            ws["A1"].font = ws["A1"].font.copy(bold=True, size=14)

            ws["A2"] = f"Период: {period_start.strftime('%d.%m.%Y')} — {period_end.strftime('%d.%m.%Y')}"

            # KPI
            ws["A4"] = "Показатель"
            ws["B4"] = "Значение"
            ws["A4"].font = ws["A4"].font.copy(bold=True)
            ws["B4"].font = ws["B4"].font.copy(bold=True)

            revenue = data["revenue"]
            ws["A5"] = "Выручка"
            ws["B5"] = revenue.get("revenue", 0)

            ws["A6"] = "Заказов оплачено"
            ws["B6"] = revenue.get("orders_count", 0)

            ws["A7"] = "Средний чек"
            ws["B7"] = revenue.get("avg_check", 0)

            # Способы оплаты
            ws["A9"] = "Способы оплаты"
            ws["A9"].font = ws["A9"].font.copy(bold=True)
            ws["B9"] = "Сумма"
            ws["B9"].font = ws["B9"].font.copy(bold=True)

            row = 10
            labels = {
                "cash": "Наличные",
                "card": "Карта",
                "transfer": "Перевод",
                "bank": "Безнал",
            }
            for method, amount in data["payment_methods"].items():
                ws[f"A{row}"] = labels.get(method, method)
                ws[f"B{row}"] = amount
                row += 1

            # Платежи
            ws[f"A{row + 1}"] = "Последние платежи"
            ws[f"A{row + 1}"].font = ws[f"A{row + 1}"].font.copy(bold=True)

            headers = ["Дата", "Заказ", "Сумма", "Способ", "Статус"]
            header_row = row + 2
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=header_row, column=col, value=header)
                cell.font = cell.font.copy(bold=True)

            row = header_row + 1
            for payment in data["payments"][:50]:  # максимум 50
                date_str = (
                    payment.paid_at.strftime("%d.%m.%Y %H:%M")
                    if payment.paid_at else "—"
                )
                ws.cell(row=row, column=1, value=date_str)
                ws.cell(row=row, column=2, value=payment.order_id)
                ws.cell(row=row, column=3, value=float(payment.amount))
                ws.cell(row=row, column=4, value=payment.method)
                ws.cell(row=row, column=5, value=payment.status)
                row += 1

            # Ширина столбцов
            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 20
            ws.column_dimensions["C"].width = 15
            ws.column_dimensions["D"].width = 15
            ws.column_dimensions["E"].width = 15

            wb.save(path)
            print(f"[ExportService] Excel сохранён: {path}")
            return True

        except Exception as e:
            print(f"[ExportService] Ошибка Excel: {e}")
            return False

    # ============================================================
    # CSV
    # ============================================================
    def _export_to_csv(self, data: Dict, path: str) -> bool:
        """Экспорт в CSV.

        Args:
            data: Данные отчёта.
            path: Путь к файлу.

        Returns:
            True если успешно.
        """
        import csv

        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f, delimiter=";")

                # Заголовок
                writer.writerow(["Финансовый отчёт CleanOps"])
                writer.writerow([
                    "Период",
                    f"{data['period_start'].strftime('%d.%m.%Y')} — "
                    f"{data['period_end'].strftime('%d.%m.%Y')}",
                ])
                writer.writerow([])

                # KPI
                writer.writerow(["Показатель", "Значение"])
                revenue = data["revenue"]
                writer.writerow(["Выручка", revenue.get("revenue", 0)])
                writer.writerow(["Заказов оплачено", revenue.get("orders_count", 0)])
                writer.writerow(["Средний чек", revenue.get("avg_check", 0)])
                writer.writerow([])

                # Способы оплаты
                writer.writerow(["Способы оплаты", "Сумма"])
                labels = {
                    "cash": "Наличные",
                    "card": "Карта",
                    "transfer": "Перевод",
                    "bank": "Безнал",
                }
                for method, amount in data["payment_methods"].items():
                    writer.writerow([labels.get(method, method), amount])
                writer.writerow([])

                # Платежи
                writer.writerow(["Дата", "Заказ", "Сумма", "Способ", "Статус"])
                for payment in data["payments"]:
                    date_str = (
                        payment.paid_at.strftime("%d.%m.%Y %H:%M")
                        if payment.paid_at else "—"
                    )
                    writer.writerow([
                        date_str,
                        payment.order_id,
                        float(payment.amount),
                        payment.method,
                        payment.status,
                    ])

            print(f"[ExportService] CSV сохранён: {path}")
            return True

        except Exception as e:
            print(f"[ExportService] Ошибка CSV: {e}")
            return False

    # ============================================================
    # PDF
    # ============================================================
    def _export_to_pdf(
        self,
        data: Dict,
        path: str,
        period_start: date,
        period_end: date,
    ) -> bool:
        """Экспорт в PDF.

        Args:
            data: Данные отчёта.
            path: Путь к файлу.
            period_start: Начало периода.
            period_end: Конец периода.

        Returns:
            True если успешно.
        """
        try:
            doc = SimpleDocTemplate(
                path,
                pagesize=A4,
                leftMargin=20 * mm,
                rightMargin=20 * mm,
                topMargin=20 * mm,
                bottomMargin=20 * mm,
            )
            styles = getSampleStyleSheet()
            story = []

            # Заголовок
            title = Paragraph("Финансовый отчёт CleanOps", styles["Title"])
            story.append(title)

            subtitle = Paragraph(
                f"Период: {period_start.strftime('%d.%m.%Y')} — "
                f"{period_end.strftime('%d.%m.%Y')}",
                styles["Normal"],
            )
            story.append(subtitle)
            story.append(Spacer(1, 10 * mm))

            # KPI
            revenue = data["revenue"]
            kpi_data = [
                ["Показатель", "Значение"],
                ["Выручка", f"{revenue.get('revenue', 0):,.0f} ₽"],
                ["Заказов оплачено", str(revenue.get("orders_count", 0))],
                ["Средний чек", f"{revenue.get('avg_check', 0):,.0f} ₽"],
            ]
            kpi_table = Table(kpi_data, colWidths=[100 * mm, 70 * mm])
            kpi_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0EA5A4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ]))
            story.append(kpi_table)
            story.append(Spacer(1, 10 * mm))

            # Способы оплаты
            story.append(Paragraph("Способы оплаты", styles["Heading2"]))
            labels = {
                "cash": "Наличные",
                "card": "Карта",
                "transfer": "Перевод",
                "bank": "Безнал",
            }
            methods_data = [["Способ", "Сумма"]]
            for method, amount in data["payment_methods"].items():
                methods_data.append([
                    labels.get(method, method),
                    f"{amount:,.0f} ₽",
                ])
            methods_table = Table(methods_data, colWidths=[100 * mm, 70 * mm])
            methods_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0EA5A4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ]))
            story.append(methods_table)

            doc.build(story)
            print(f"[ExportService] PDF сохранён: {path}")
            return True

        except Exception as e:
            print(f"[ExportService] Ошибка PDF: {e}")
            return False
